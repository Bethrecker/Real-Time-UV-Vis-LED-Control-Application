from bootstrap import configure
configure()
from dataclasses import replace, asdict
from concurrent.futures import Future
import csv
import json
from pathlib import Path
import queue
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from acquisition import SimulationSpectrometer, SpectrometerClient, acquire_capture
from control import ExperimentController
from led import LEDService, MightexAdapter, OutputApproval, SimulationLED
from models import Capture, RunSettings, TimingBudget, TimingOverrun
from processing import average, boxcar, process, usable_reference
from session import Session
from led_session import LEDConnection
from spectrometer import Settings, Spectrum
from storage import RunStore, combined_export


ROOT = Path(__file__).resolve().parent


def capture(value, settings=Settings(10000, 1, 0), serial="SIM-QEP06321"):
    scan = tuple([value]*100)
    spectrum = Spectrum(tuple(range(200, 300)), scan, settings, serial, "start", "finish", 0.01)
    return Capture(spectrum, (scan,), {"scan_calls": []})


class ProcessingTests(unittest.TestCase):
    def test_arithmetic_mean_and_boxcar_edges(self):
        self.assertEqual(average([(1, 2, 3), (3, 4, 5)]), (2, 3, 4))
        self.assertEqual(boxcar((0, 6, 0), 1), (3, 2, 3))

    def test_blocked_reference_is_refused(self):
        self.assertFalse(usable_reference(capture(1500), capture(1510), 200000).usable)

    def test_known_absorbance_and_invalid_sample(self):
        dark, blank = capture(1500), capture(10000)
        result = process(capture(5750), dark, blank, 200000)
        self.assertAlmostEqual(result.absorbance[0], 0.301029995664)
        bad = process(capture(1500), dark, blank, 200000)
        self.assertTrue(all(v is None for v in bad.absorbance))
        self.assertTrue(all(v == "nonpositive_sample_minus_dark" for v in bad.flags))

    def test_peak_saturation_not_hidden_by_averaging(self):
        s = Settings(10000, 2, 2)
        dark, blank = capture(1500, s), capture(10000, s)
        first, second = [5750]*100, [5750]*100
        first[40] = 200000
        source = Capture(replace(capture(5750, s).spectrum, intensities=average([first, second])), (tuple(first), tuple(second)), {})
        result = process(source, dark, blank, 200000)
        self.assertEqual(result.flags[40], "saturated_sdk_intensity")
        self.assertTrue(all(result.flags[i] == "valid" for i in (38,39,41,42)))
        self.assertEqual(result.flags[37], "valid")

    def test_mismatched_background_settings_rejected(self):
        self.assertFalse(usable_reference(capture(1500), capture(10000, Settings(20000, 1, 0)), 200000).usable)


class BudgetTests(unittest.TestCase):
    def test_defaults_fit_and_short_off_rejected(self):
        settings = RunSettings()
        settings.validate()
        self.assertAlmostEqual(settings.timing.collection_window(settings.acquisition), 0.450)
        self.assertEqual(settings.timing, TimingBudget(0.020, 0.020, 0.020, 0.200, 0.500))
        self.assertEqual(settings.duration_s, 200)
        with self.assertRaises(ValueError):
            replace(settings, off_s=0.690).validate()
        replace(settings, off_s=0.691).validate()
        with self.assertRaises(ValueError):
            replace(settings, off_s=0.4).validate()

    def test_invalid_numbers_and_pulse_quantization(self):
        for value in [replace(RunSettings(), on_s=float("nan")), replace(RunSettings(), cycles=0), replace(RunSettings(), on_s=0.000021)]:
            with self.assertRaises(ValueError):
                value.validate()

    def test_slow_read_causes_overrun(self):
        source = SimulationSpectrometer()
        source.apply(Settings(10000, 1, 0))
        original = source.single
        def slow():
            time.sleep(0.06)
            return original()
        source.single = slow
        with self.assertRaises(TimingOverrun):
            acquire_capture(source, TimingBudget(), threading.Event())


class StorageTests(unittest.TestCase):
    def test_incremental_recovery_and_excel_integrity(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            settings = RunSettings(cycles=2, acquisition=Settings(10000, 1, 0))
            dark, blank, sample = capture(1500), capture(10000), capture(5750)
            store = RunStore(temp, "test", settings, dark, blank, {"simulation": True})
            store.save_cycle(1, sample, process(sample, dark, blank, 200000), {"provenance": "host only"}).result(3)
            store.finish("stopped").result(3)
            directory = store.directory
            store.close()
            self.assertTrue((directory / "cycle_000001.json").exists())
            self.assertEqual(json.loads((directory / "manifest.json").read_text())["status"], "stopped")
            combined_export(directory)
            with (directory / "combined_absorbance.csv").open(newline="") as stream:
                rows = list(csv.reader(stream))
            self.assertEqual(rows[0][0], "wavelength_nm")
            self.assertAlmostEqual(float(rows[1][1]), 0.301029995664)
            from openpyxl import load_workbook
            workbook = load_workbook(directory / "combined.xlsx", read_only=True)
            self.assertEqual(workbook["Absorbance"]["A1"].value, "wavelength_nm")
            self.assertAlmostEqual(workbook["Absorbance"]["B2"].value, float(rows[1][1]))
            self.assertIn("SDK intensities", workbook.sheetnames)
            workbook.close()


class LEDTests(unittest.TestCase):
    def test_approval_required_and_no_guess(self):
        for approval in [OutputApproval(0, 100, True), OutputApproval(1, 100, False), OutputApproval(1, 0, True)]:
            with self.assertRaises(ValueError):
                approval.validate()

    def test_documented_pulse_and_explicit_disable_before_close(self):
        class FakeDLL:
            calls = []
            def __getattr__(self, name):
                def call(*args):
                    self.calls.append((name, args))
                    return 0
                return call
        adapter = MightexAdapter(OutputApproval(2, 100, True))
        adapter.dll, adapter.handle = FakeDLL(), 0
        adapter.prepare(1)
        adapter.on(threading.Event())
        self.assertEqual(adapter.close(), [])
        calls = adapter.dll.calls
        detail = next(args for name, args in calls if name.endswith("SetPulseDetail"))
        self.assertEqual(detail, (0, 2, 0, 20, 1000000, 20, 0, 100, 0))
        self.assertTrue(calls[-2][0].endswith("SetMode"))
        self.assertEqual(calls[-2][1][-1], 0)
        self.assertTrue(calls[-1][0].endswith("CloseDevice"))
        self.assertFalse(any("Reset" in name or "Store" in name for name, args in calls))

    def test_stop_inhibits_queued_pulse(self):
        adapter = SimulationLED()
        service = LEDService(adapter)
        service.arm()
        gate = threading.Event()
        blocker = service.executor.submit(gate.wait)
        pulse = service.on()
        off = service.stop()
        gate.set()
        with self.assertRaises(RuntimeError):
            pulse.result(1)
        off.result(1)
        self.assertFalse(adapter.enabled)
        self.assertEqual(service.close(), [])

    def test_disable_failure_reported(self):
        adapter = SimulationLED()
        adapter.off = lambda: (_ for _ in ()).throw(RuntimeError("injected off failure"))
        service = LEDService(adapter)
        self.assertTrue(service.close())


class ControllerTests(unittest.TestCase):
    def test_blocked_acquisition_does_not_delay_stop_off(self):
        stop = threading.Event()
        blocked = threading.Event()
        release = threading.Event()
        settings = RunSettings(on_s=0.02, off_s=0.5, cycles=1, acquisition=Settings(10000, 1, 0))
        dark, blank = capture(1500), capture(10000)
        adapter = SimulationLED()
        led = LEDService(adapter)
        class Spec:
            identity = SimulationSpectrometer().identity
            def capture(self, *args):
                blocked.set()
                release.wait(2)
                from models import Stopped
                raise Stopped()
        class Store:
            def finish(self, *args):
                future = Future();future.set_result(None);return future
        controller = ExperimentController(led, Spec(), Store(), stop, lambda *args: None)
        worker = threading.Thread(target=controller.run, args=(settings, dark, blank))
        worker.start()
        self.assertTrue(blocked.wait(2))
        before = len(adapter.calls)
        stop.set()
        led.stop().result(0.5)
        self.assertTrue(worker.is_alive())
        self.assertFalse(adapter.enabled)
        self.assertEqual(adapter.calls[before][0], "off")
        release.set();worker.join(2);led.close()

    def test_overrun_prevents_next_pulse_and_always_off(self):
        settings = RunSettings(on_s=0.02, off_s=0.5, cycles=2, acquisition=Settings(10000, 1, 0))
        adapter, stop = SimulationLED(), threading.Event()
        led = LEDService(adapter)
        class Spec:
            identity = SimulationSpectrometer().identity
            def capture(self, *args):
                raise TimingOverrun("injected blocked-read deadline")
        class Store:
            def finish(self, *args):
                f=Future();f.set_result(None);return f
        controller = ExperimentController(led, Spec(), Store(), stop, lambda *args: None)
        result = controller.run(settings, capture(1500), capture(10000))
        self.assertEqual(result[0], "error")
        self.assertEqual(sum(name == "on" for name, stamp in adapter.calls), 0)
        self.assertFalse(adapter.enabled)
        led.close()


class SimulationIntegrationTests(unittest.TestCase):
    def test_process_fresh_capture_and_full_two_cycle_session(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as temp, patch("storage.DEFAULT_RUN_CACHE", Path(temp)/"recovery"):
            settings = RunSettings(on_s=0.04, off_s=0.6, cycles=2, acquisition=Settings(10000, 1, 0))
            events = queue.Queue()
            connection = LEDConnection(True, events)
            connection.start()
            self.assertTrue(connection.ready.wait(2))
            self.addCleanup(connection.join, 5)
            self.addCleanup(connection.halt)
            session = Session(True, ROOT, "unused", None, settings, events, led_connection=connection)
            session.start()
            def await_kind(wanted, timeout=10):
                until = time.monotonic() + timeout
                while time.monotonic() < until:
                    kind, value = events.get(timeout=max(0.01, until-time.monotonic()))
                    if kind == "error":
                        self.fail(value)
                    if kind == wanted:
                        return value
                self.fail("Timed out: " + wanted)
            await_kind("connected")
            session.commands.put(("preview", (False, "blocked")))
            session.commands.put(("dark", None));await_kind("dark")
            session.commands.put(("preview", (False, "blank")))
            session.commands.put(("reference", None));await_kind("reference")
            session.commands.put(("start", (temp, "integration", True)))
            directory = Path(await_kind("run_folder"))
            result = await_kind("run_finished")
            self.assertEqual(result, ("completed", None, 2))
            await_kind("run_ready")
            self.assertTrue(session.is_alive())
            self.assertTrue(connection.ready.is_set())
            self.assertFalse(session.experiment_started)
            self.assertFalse(connection.reserved)
            session.halt()
            self.assertEqual(await_kind("disconnected"), [])
            session.join(2)
            records = list(session.store.working_directory.glob("cycle_*.json"))
            self.assertEqual(len(records), 3)
            self.assertEqual([p.name for p in directory.iterdir()], ["combined_absorbance.xlsx"])
            baseline = json.loads((session.store.working_directory / "cycle_000000.json").read_text())
            self.assertEqual(baseline["timing"]["irradiation_elapsed_s"], 0)
            record = json.loads((session.store.working_directory / "cycle_000001.json").read_text())
            timing = record["timing"]
            center = (timing["collection_target_start_monotonic_s"] + timing["collection_target_end_monotonic_s"])/2
            self.assertAlmostEqual(center, (timing["off_start_host_ack_monotonic_s"] + timing["off_end_target_monotonic_s"])/2)
            self.assertFalse(record["capture"]["provenance"]["optical_timing_verified"])
            self.assertFalse(session.led.adapter.enabled)


class GUITests(unittest.TestCase):
    def test_complete_gui_simulation_workflow(self):
        import tkinter as tk
        from gui import Application
        root = tk.Tk();root.withdraw()
        from test_support import make_application
        app = make_application(root)
        with tempfile.TemporaryDirectory(dir=ROOT) as temp, patch("storage.DEFAULT_RUN_CACHE", Path(temp)/"recovery"):
            try:
                for key, value in {"on": "0.04", "off": "0.6", "cycles": "2", "scans": "1", "boxcar": "0", "folder": temp, "lateness": "100"}.items():
                    app.variables[key].set(value)
                def wait_for(predicate, seconds=10):
                    until = time.monotonic()+seconds
                    while time.monotonic() < until:
                        root.update()
                        if predicate():
                            return
                        time.sleep(0.01)
                    self.fail("GUI state transition timed out: " + app.warning.get())
                app.connect();wait_for(lambda: app.identity is not None and not app.pending)
                app.connect_led();wait_for(lambda: app.led_session.ready.is_set() and not app.led_pending)
                app.capture("dark");wait_for(lambda: app.dark is not None and not app.pending)
                # Still blocked: this reference must fail the Start gate.
                app.capture("reference");wait_for(lambda: app.reference is not None and not app.pending)
                self.assertEqual(str(app.start_button["state"]), "disabled")
                app.probe_phase.set("blank");app.toggle_preview()
                app.capture("reference");wait_for(lambda: not app.pending)
                self.assertEqual(str(app.start_button["state"]), "normal")
                app.start();wait_for(lambda: not app.running and not app.pending, 15)
                self.assertEqual(app.completed, 2, app.status.get()+" "+app.warning.get())
                self.assertEqual(len(app.abs_lines), 3)
                self.assertEqual(app.abs_lines[-1].get_linewidth(), 2)
                self.assertTrue((Path(app.run_path)/"combined_absorbance.xlsx").is_file())
                session, connection = app.session, app.led_session
                self.assertTrue(session.is_alive())
                self.assertTrue(connection.ready.is_set())
                self.assertFalse(connection.led.adapter.enabled)
                previous_output = Path(app.run_path)/"combined_absorbance.xlsx"
                previous_bytes = previous_output.read_bytes()
                with patch("gui.simpledialog.askstring", return_value="next_slide"):
                    app.new_run()
                self.assertEqual(app.abs_lines, [])
                self.assertIs(app.session, session)
                self.assertIs(app.led_session, connection)
                self.assertEqual(str(app.background_button["state"]), "normal")
                app.background_blank()
                wait_for(lambda: not app.pending)
                app.start()
                wait_for(lambda: not app.running and not app.pending, 15)
                self.assertEqual(app.completed, 2, app.warning.get())
                self.assertIs(app.session, session)
                self.assertIs(app.led_session, connection)
                self.assertTrue(connection.ready.is_set())
                self.assertFalse(connection.led.adapter.enabled)
                self.assertNotEqual(Path(app.run_path)/"combined_absorbance.xlsx", previous_output)
                self.assertEqual(previous_output.read_bytes(), previous_bytes)
                app.abs_canvas.figure.savefig(ROOT/"verification_absorbance.png")
            finally:
                if app.session:
                    app.session.halt();app.session.join(5)
                if app.led_session:
                    app.led_session.halt();app.led_session.join(5)
                app.export_worker.shutdown();app.destroy_window()

    def test_gui_builds_and_gates_reference_and_real_approval(self):
        import tkinter as tk
        from gui import Application
        root = tk.Tk();root.withdraw()
        from test_support import make_application
        app = make_application(root)
        try:
            root.update()
            self.assertEqual(str(app.start_button["state"]), "disabled")
            app.update()
            self.assertEqual(str(app.connect_button["state"]), "normal")
            self.assertEqual(str(app.led_connect_button["state"]), "normal")
            self.assertEqual(str(app.led_test_button["state"]), "disabled")
            app.variables["channel"].set("1")
            app.variables["level"].set("10")
            app.confirmed.set(True);app.update()
            self.assertEqual(str(app.connect_button["state"]), "normal")
            self.assertEqual(str(app.start_button["state"]), "disabled")
        finally:
            app.export_worker.shutdown();app.destroy_window()


if __name__ == "__main__":
    unittest.main()
