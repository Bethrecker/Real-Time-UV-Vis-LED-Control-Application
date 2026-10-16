"""Version 2.1 new-run and Excel-only export checks; no hardware."""
import json
import math
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from models import RunSettings
from processing import process
from storage import RunStore, copy_file_atomically, recover_absorbance_excel
from test_application import capture
import test_run_timing as gui_fixture
from gui import VERSION_CREDIT

ROOT = Path(__file__).resolve().parent


class WorkbookCopyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT)
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name)/"local.xlsx"
        self.target = Path(self.temp.name)/"output.xlsx"
        self.source.write_bytes(b"new workbook")
        self.target.write_bytes(b"previous workbook")

    def test_transient_replace_lock_retries_and_copies_complete_file(self):
        import os
        replace = os.replace
        with patch("storage.os.replace", side_effect=[PermissionError("sync lock"), None]) as mocked, patch("storage.time.sleep") as sleep:
            # The successful attempt must actually replace the file.
            def locked_once(source, target):
                if mocked.call_count == 1:
                    raise PermissionError("sync lock")
                return replace(source, target)
            mocked.side_effect = locked_once
            copy_file_atomically(self.source, self.target)
        self.assertEqual(mocked.call_count, 2)
        sleep.assert_called_once_with(0.1)
        self.assertEqual(self.target.read_bytes(), self.source.read_bytes())
        self.assertFalse(self.target.with_suffix(".xlsx.tmp").exists())

    def test_transient_temporary_file_lock_retries(self):
        original_open = Path.open
        locked = [True]
        def open_file(path, mode="r", *args, **kwargs):
            if mode == "wb" and locked[0]:
                locked[0] = False
                raise PermissionError("temporary file sync lock")
            return original_open(path, mode, *args, **kwargs)
        with patch.object(Path, "open", open_file), patch("storage.time.sleep") as sleep:
            copy_file_atomically(self.source, self.target)
        sleep.assert_called_once_with(0.1)
        self.assertEqual(self.target.read_bytes(), self.source.read_bytes())

    def test_persistent_lock_preserves_old_destination_and_local_workbook(self):
        with patch("storage.os.replace", side_effect=PermissionError("open workbook")) as replace, patch("storage.time.sleep") as sleep:
            with self.assertRaisesRegex(PermissionError, "Re-export a saved/interrupted run") as error:
                copy_file_atomically(self.source, self.target)
        self.assertEqual(replace.call_count, 5)
        self.assertAlmostEqual(sum(call.args[0] for call in sleep.call_args_list), 1.5)
        self.assertIn(str(self.source), str(error.exception))
        self.assertEqual(self.target.read_bytes(), b"previous workbook")
        self.assertEqual(self.source.read_bytes(), b"new workbook")
        self.assertFalse(self.target.with_suffix(".xlsx.tmp").exists())

    def test_other_io_errors_are_not_retried(self):
        with patch("storage.os.replace", side_effect=OSError("disk full")) as replace, patch("storage.time.sleep") as sleep:
            with self.assertRaisesRegex(OSError, "disk full"):
                copy_file_atomically(self.source, self.target)
        replace.assert_called_once()
        sleep.assert_not_called()
        self.assertEqual(self.target.read_bytes(), b"previous workbook")


class ExcelOnlyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT)
        self.base = Path(self.temp.name)
        self.output = self.base/"output"
        self.output.mkdir()
        self.cache_patch = patch("storage.DEFAULT_RUN_CACHE", self.base/"recovery")
        self.cache_patch.start()
        self.settings = RunSettings(cycles=3)
        self.dark, self.reference, self.sample = capture(1500), capture(10000), capture(5750)

    def tearDown(self):
        self.cache_patch.stop()
        self.temp.cleanup()

    def store(self, on_error=None):
        return RunStore(self.output, "excel_only", self.settings, self.dark, self.reference, {},
                        on_error=on_error, combined_excel_only=True)

    def save(self, store, cycle):
        return store.save_cycle(cycle, self.sample, process(self.sample, self.dark, self.reference, 200000), {})

    def check_workbook(self, store, cycles):
        self.assertEqual([p.name for p in store.directory.iterdir()], ["combined_absorbance.xlsx"])
        book = load_workbook(store.directory/"combined_absorbance.xlsx", read_only=True)
        try:
            self.assertEqual(book.sheetnames, ["Absorbance"])
            sheet = book["Absorbance"]
            rows = list(sheet.iter_rows(values_only=True))
            self.assertEqual(rows[0], tuple(["wavelength_nm"]+[f"cycle_{i:06d}_absorbance_AU" for i in range(1,cycles+1)]))
            self.assertEqual(rows[1][0], 200)
            for value in rows[1][1:]:
                self.assertAlmostEqual(value, math.log10(2))
        finally:
            book.close()

    def test_only_combined_excel_is_visible_and_updated_incrementally(self):
        store = self.store()
        try:
            self.assertNotEqual(store.directory, store.working_directory)
            store._drain_mirror()
            self.check_workbook(store, 0)
            for cycle in (1,2,3):
                self.save(store, cycle).result(2)
                store._drain_mirror()
                self.check_workbook(store, cycle)
            store.finish("completed").result(2)
            store.export().result(5)
            self.check_workbook(store, 3)
            private = store.working_directory
            self.assertEqual(len(list(private.glob("cycle_*.json"))), 3)
            self.assertTrue((private/"dark.json").is_file())
            self.assertTrue((private/"reference.json").is_file())
            self.assertEqual(json.loads((private/"manifest.json").read_text())["status"], "completed")
        finally:
            store.close()

    def test_excel_only_cycle_recovers_without_csv_write(self):
        store = self.store()
        try:
            # An unnecessary CSV flush must not delay a durable JSON commit.
            with patch.object(store, "_write_capture_csv", side_effect=OSError("CSV unavailable")):
                self.save(store, 1).result(2)
            self.assertTrue((store.working_directory/"cycle_000001.json").is_file())
            self.assertFalse((store.working_directory/"cycle_000001.csv").exists())
            store.finish("stopped").result(2)
            store.export().result(5)
            self.check_workbook(store, 1)
            recover_absorbance_excel(store.working_directory)
            self.check_workbook(store, 1)
        finally:
            store.close()

    def test_slow_excel_export_does_not_block_durable_cycle_saving(self):
        entered, release = threading.Event(), threading.Event()
        from storage import copy_file_atomically
        def slow_copy(source, target):
            entered.set()
            if not release.wait(8):
                raise TimeoutError("test export gate")
            copy_file_atomically(source, target)
        with patch("storage.copy_file_atomically", slow_copy):
            store = self.store()
            try:
                self.assertTrue(entered.wait(2))
                for cycle in (1,2,3):
                    self.save(store, cycle).result(1)
                self.assertEqual(len(list(store.working_directory.glob("cycle_*.json"))), 3)
                store.finish("stopped").result(1)
                release.set()
                store.export().result(5)
                self.check_workbook(store, 3)
            finally:
                release.set()
                store.close()

    def test_export_failure_reports_error_and_recovery_restores_all_completed_data(self):
        failed = threading.Event()
        store = self.store(lambda message: failed.set())
        try:
            self.save(store, 1).result(2)
            store._drain_mirror()
            with patch("storage.copy_file_atomically", side_effect=PermissionError("workbook is open")):
                self.save(store, 2).result(2)
                self.assertTrue(failed.wait(2))
                with self.assertRaisesRegex(OSError, "workbook is open"):
                    store.check_health()
                store.finish("error", "export failed").result(2)
                with self.assertRaisesRegex(OSError, "workbook is open"):
                    store.export().result(5)
            self.assertTrue((store.working_directory/"cycle_000002.json").exists())
            final = recover_absorbance_excel(store.directory)
            self.assertEqual(final, str(store.directory/"combined_absorbance.xlsx"))
            self.check_workbook(store, 2)
        finally:
            store.close()

    def test_export_backlog_stops_saving_requests_without_losing_recovery_record(self):
        release, entered = threading.Event(), threading.Event()
        def blocked_copy(*args):
            entered.set()
            release.wait(8)
        with patch("storage.copy_file_atomically", blocked_copy):
            store = self.store()
            try:
                self.assertTrue(entered.wait(2))
                for cycle in range(1,16):
                    self.save(store, cycle).result(1)
                with self.assertRaisesRegex(OSError, "fell behind by 16 batches"):
                    self.save(store, 16).result(1)
                self.assertTrue((store.working_directory/"cycle_000016.json").exists())
            finally:
                release.set()
                store.close()

    def test_excel_column_limit_rejected_before_creating_run(self):
        from dataclasses import replace
        with self.assertRaisesRegex(ValueError, "16,382"):
            RunStore(self.output, "too_many", replace(self.settings, cycles=16383), self.dark, self.reference, {}, combined_excel_only=True)
        self.assertEqual(list(self.output.iterdir()), [])


class NewRunTests(unittest.TestCase):
    setUp = gui_fixture.LiveAbsorbanceGUITests.setUp
    tearDown = gui_fixture.LiveAbsorbanceGUITests.tearDown

    def test_credit_and_controls_are_in_requested_positions(self):
        app = self.app
        self.assertEqual(app.credit_label.cget("text"), VERSION_CREDIT)
        self.assertEqual(app.status.get(), "")
        self.assertEqual(app.warning.get(), "")
        self.assertIs(app.experiment_controls.master, app.led_connect_button.master)
        self.assertEqual(app.experiment_controls.pack_info()["side"], "right")
        self.assertEqual(app.new_run_button.pack_info()["side"], "right")
        self.assertFalse(hasattr(app,"excel_button"))

    def test_new_run_retains_settings_and_backgrounds_clears_display_and_preserves_saved_files(self):
        app = self.app
        app.variables["integration"].set("25")
        previous = {key:value.get() for key,value in app.variables.items()}
        confirmed = app.confirmed.get()
        app.completed = app.active_cycle = 2
        app.requested_cycles = 3
        app.progress.configure(value=2)
        dark, reference = capture(1500), capture(10000)
        app.dark, app.reference, app.latest = dark, reference, capture(5750)
        line, = app.abs_axes.plot([200,201],[.3,.4])
        app.abs_lines.append(line)
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            file = Path(temp)/"previous.xlsx"
            file.write_bytes(b"preserved saved experiment")
            app.run_path = temp
            with patch("gui.simpledialog.askstring", return_value="next_sample") as dialog:
                app.new_run()
            dialog.assert_called_once()
            self.assertEqual(file.read_bytes(), b"preserved saved experiment")
        self.assertEqual(app.variables["name"].get(),"next_sample")
        self.assertEqual({key:value.get() for key,value in app.variables.items() if key!="name"},
                         {key:value for key,value in previous.items() if key!="name"})
        self.assertEqual(app.confirmed.get(), confirmed)
        self.assertIs(app.dark,dark)
        self.assertIs(app.reference,reference)
        self.assertIsNone(app.latest)
        self.assertIsNone(app.run_path)
        self.assertEqual(app.completed,0)
        self.assertEqual(len(app.abs_axes.lines),0)
        self.assertEqual(app.abs_lines,[])

    def test_cancel_and_unchanged_name_do_not_clear_until_different_name_confirmed(self):
        app = self.app
        app.completed = 2
        with patch("gui.simpledialog.askstring",return_value=None):
            app.new_run()
        self.assertEqual(app.completed,2)
        previous = app.variables["name"].get()
        with patch("gui.simpledialog.askstring",side_effect=[previous, "new_sample"]) as dialog:
            app.new_run()
        self.assertEqual(dialog.call_count,2)
        self.errors.assert_called_once()
        self.assertEqual(app.completed,0)

    def test_new_run_does_not_interrupt_acquisition_led_testing_shutdown_or_export(self):
        app = self.app
        with patch("gui.simpledialog.askstring") as dialog:
            for flag in ("running","pending","led_testing","manual_export_pending"):
                setattr(app,flag,True)
                app.new_run()
                setattr(app,flag,False)
            app.stop_time = 1
            app.new_run()
            app.stop_time = None
            dialog.assert_not_called()


if __name__ == "__main__":
    unittest.main()
