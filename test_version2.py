"""Version 2 SDK acquisition contracts. All devices are explicit fakes."""
from bootstrap import configure
configure()
from dataclasses import replace
import csv
import json
import math
from pathlib import Path
import queue
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch
from acquisition import FreshOceanSpectrometer, acquire_capture, acquire_preview
from led import OutputApproval
from led_session import LEDConnection
from models import Capture, PreviewProcessing, RunSettings, TimingBudget, SDK_PROCESSING
from processing import process
from session import Session
from spectrometer import Identity, Settings, Spectrum
from test_spectrum_update import VendorDevice
from test_support import SilentSound

ROOT = Path(__file__).resolve().parent


def vendor_adapter():
    adapter = FreshOceanSpectrometer()
    adapter.identity = Identity("QEP06321", "QE-PRO", (1, 30, 4), 10000, 1600000000, 1, 200000)
    adapter.wavelengths = tuple(range(200, 300))
    adapter.buffer_policy = {"mode": "disabled_enable_flag"}
    adapter.device = VendorDevice()
    adapter.device.result = [5750.0] * 100
    return adapter


class Version2Tests(unittest.TestCase):
    def test_gui_has_no_simulation_choice_and_connects_only_on_request(self):
        import tkinter as tk
        from gui import Application
        root = tk.Tk()
        root.withdraw()
        spec_factory, led_factory = Mock(), Mock()
        app = Application(root, spec_factory, led_factory, sound_factory=SilentSound)
        try:
            spec_factory.assert_not_called()
            led_factory.assert_not_called()
            for name in ("simulation", "mode_button", "phase_selector", "probe_phase",
                         "preview_electric_dark", "preview_nonlinearity", "electric_dark_button", "nonlinearity_button"):
                self.assertFalse(hasattr(app, name), name)
            self.assertIn("REAL HARDWARE", app.mode_status.get())
            self.assertEqual(app.preview_settings().timing, TimingBudget())
            self.assertEqual(app.preview_settings().preview_processing, SDK_PROCESSING)
            self.assertEqual(app.variables["channel"].get(), "")
            self.assertEqual(app.variables["level"].get(), "")
            app.connect()
            self.assertIs(spec_factory.call_args.args[0], False)
            led_factory.assert_not_called()
            app.connect_led()
            self.assertIs(led_factory.call_args.args[0], False)
        finally:
            app.export_worker.shutdown()
            app.destroy_window()

    def test_all_capture_modes_preserve_sdk_result_and_user_settings(self):
        adapter = vendor_adapter()
        settings = Settings(25000, 3, 3)
        adapter.apply(settings)
        budget = TimingBudget()
        cancel = threading.Event()
        preview = acquire_preview(adapter, budget, PreviewProcessing(True, True), cancel)
        self.assertEqual(preview.spectrum.settings, settings)
        self.assertEqual(preview.provenance["preview_processing"], {"electric_dark": True, "nonlinearity": True})
        for _role in ("dark", "reference", "sample"):
            capture = acquire_capture(adapter, budget, cancel)
            self.assertEqual(capture.spectrum.settings, settings)
            self.assertEqual(capture.spectrum.intensities, tuple(adapter.device.result))
            self.assertEqual(len(capture.provenance["scan_calls"]), 1)
            self.assertFalse(capture.provenance["host_averaging_applied"])
            self.assertFalse(capture.provenance["host_boxcar_applied"])
            self.assertFalse(capture.provenance["raw_source_scans_available"])
            self.assertTrue(capture.provenance["sdk_settings_readback_verified"])
            self.assertTrue(capture.provenance["electric_dark"])
            self.assertTrue(capture.provenance["nonlinearity"])
            self.assertFalse(capture.provenance["averaging_behavior_verified"])
        self.assertEqual(adapter.device.reads, 4)
        self.assertTrue(adapter.device.values["electric_dark_correction_usage"])
        self.assertTrue(adapter.device.values["nonlinearity_correction_usage"])
        self.assertFalse(any(not value for name, value in adapter.device.calls if name.endswith("correction_usage")))
        self.assertTrue(all(value == 3 for name, value in adapter.device.calls
                            if name in ("set_scans_to_average", "set_boxcar_width")))
        self.assertTrue(all(value == 25000 for name, value in adapter.device.calls if name == "set_integration_time"))

    def test_changed_sdk_settings_reject_spectrum_before_read(self):
        for field, value in (("integration_time", 20000), ("scans_to_average", 1), ("boxcar_width", 0)):
            adapter = vendor_adapter()
            adapter.apply(Settings(10000, 3, 2))
            adapter.device.values[field] = value
            with self.assertRaisesRegex(RuntimeError, "SDK acquisition settings changed"):
                adapter.single()
            self.assertEqual(adapter.device.reads, 0)

    def test_changed_sdk_correction_rejects_spectrum_before_read(self):
        for field in ("electric_dark_correction_usage", "nonlinearity_correction_usage"):
            adapter = vendor_adapter()
            adapter.apply(Settings(10000, 3, 2))
            adapter.device.values[field] = False
            with self.assertRaisesRegex(RuntimeError, "SDK correction settings changed"):
                adapter.single()
            self.assertEqual(adapter.device.reads, 0)

    def test_absorbance_uses_sdk_pixels_without_second_boxcar(self):
        settings = Settings(10000, 3, 2)
        def capture(values):
            spectrum = Spectrum(tuple(range(200, 300)), tuple(values), settings, "QEP06321", "start", "end", 0.01)
            return Capture(spectrum, (tuple(values),), {"raw_source_scans_available": False})
        sample = [5750.0] * 100
        sample[50] = 3200.0
        result = process(capture(sample), capture([1500.0]*100), capture([10000.0]*100), 200000)
        self.assertAlmostEqual(result.absorbance[50], -math.log10(1700/8500))
        for i in (48, 49, 51, 52):
            self.assertAlmostEqual(result.absorbance[i], -math.log10(0.5))

    def test_sdk_read_budget_includes_selected_averaging_exposures(self):
        settings = Settings(25000, 4, 2)
        budget = TimingBudget()
        self.assertAlmostEqual(budget.sdk_read_budget(settings), 0.260)
        self.assertAlmostEqual(budget.collection_window(settings), 0.780)
        with self.assertRaisesRegex(ValueError, "Off-period too short"):
            replace(RunSettings(), off_s=0.90, acquisition=settings).validate()
        adapter = vendor_adapter()
        adapter.apply(settings)
        adapter.device.delay = 0.120
        capture = acquire_capture(adapter, budget, threading.Event())
        self.assertEqual(capture.spectrum.settings, settings)

    def test_complete_session_saves_vendor_settings_for_backgrounds_and_cycles(self):
        adapter = vendor_adapter()
        read_settings = []
        original_read = adapter.device.get_formatted_spectrum
        def recorded_read():
            v = adapter.device.values
            read_settings.append(Settings(v["integration_time"], v["scans_to_average"], v["boxcar_width"]))
            return original_read()
        adapter.device.get_formatted_spectrum = recorded_read

        class FakeClient:
            identity = adapter.identity
            def __init__(self, simulation, *_args):
                if simulation:
                    raise AssertionError("Real session must request the real adapter path")
                self.cancel = threading.Event()
                self.closed = False
            def apply(self, settings):
                return adapter.apply(settings)
            def configure_preview(self, processing):
                return adapter.configure_preview(processing)
            def capture(self, budget, phase="sample", cycle=0, deadline=None, interrupt=None):
                adapter.device.result = [1500.0 if phase == "blocked" else 10000.0 if phase == "blank" else 5750.0] * 100
                return acquire_capture(adapter, budget, interrupt or self.cancel, deadline, cycle > 0)
            def close(self):
                self.closed = True
                return []

        with tempfile.TemporaryDirectory(dir=ROOT) as temp, patch("session.SpectrometerClient", FakeClient), patch("storage.DEFAULT_RUN_CACHE", Path(temp)/"recovery"):
            settings = RunSettings(on_s=0.02, off_s=1.2, cycles=3, acquisition=Settings(25000, 3, 3),
                                   timing=TimingBudget(lateness_s=0.1))
            events = queue.Queue()
            led = LEDConnection(True, events)  # explicit fake LED; no DLL load
            led.start()
            self.assertTrue(led.ready.wait(2))
            session = Session(False, ROOT, "QEP06321", OutputApproval(1, 100, True), settings, events, led_connection=led)
            session.start()
            def await_kind(wanted):
                until = time.monotonic() + 15
                while time.monotonic() < until:
                    kind, value = events.get(timeout=max(0.01, until-time.monotonic()))
                    if kind == "error":
                        self.fail(value)
                    if kind == wanted:
                        return value
                self.fail("Timed out: " + wanted)
            try:
                await_kind("connected")
                session.commands.put(("preview", (False, "blocked")))
                session.commands.put(("dark", None))
                await_kind("dark")
                session.commands.put(("preview", (False, "blank")))
                session.commands.put(("reference", None))
                await_kind("reference")
                session.commands.put(("start", (temp, "v2_sdk_contract", True)))
                directory = Path(await_kind("run_folder"))
                self.assertEqual(await_kind("run_finished"), ("completed", None, 3))
                await_kind("run_ready")
                self.assertFalse(session.spec.closed)
                self.assertTrue(led.ready.is_set())
                self.assertFalse(led.reserved)
                session.halt()
                self.assertEqual(await_kind("disconnected"), [])
                self.assertTrue(session.spec.closed)
                self.assertFalse(led.led.adapter.enabled)
                self.assertEqual(read_settings, [settings.acquisition] * 6)
                private = session.store.working_directory
                manifest = json.loads((private / "manifest.json").read_text())
                self.assertEqual(manifest["format_version"], 2)
                for name, expected in (("dark", 1500), ("reference", 10000), ("cycle_000000", 5750), ("cycle_000001", 5750),
                                       ("cycle_000002", 5750), ("cycle_000003", 5750)):
                    record = json.loads((private / f"{name}.json").read_text())
                    capture = record.get("capture", record)
                    self.assertEqual(capture["spectrum"]["settings"], {"integration_us": 25000, "scans": 3, "boxcar": 3})
                    self.assertEqual(capture["spectrum"]["intensities"], [expected]*100)
                    self.assertFalse(capture["provenance"]["host_boxcar_applied"])
                    self.assertTrue(capture["provenance"]["electric_dark"])
                    self.assertTrue(capture["provenance"]["nonlinearity"])
                    if name in ("dark", "reference"):
                        with (private / f"{name}.csv").open(newline="") as stream:
                            header = next(csv.reader(stream))
                        self.assertEqual(header[:2], ["wavelength_nm", "sdk_intensity_counts"])
                        self.assertFalse(any("raw_scan" in value for value in header))
                    else:
                        self.assertFalse((private / f"{name}.csv").exists())
                from openpyxl import load_workbook
                self.assertEqual([p.name for p in directory.iterdir()], ["combined_absorbance.xlsx"])
                workbook = load_workbook(directory / "combined_absorbance.xlsx", read_only=True)
                try:
                    self.assertEqual(workbook.sheetnames, ["Absorbance"])
                    self.assertEqual(workbook["Absorbance"]["B1"].value, "time_0_no_irradiation_absorbance_AU")
                    self.assertAlmostEqual(workbook["Absorbance"]["B2"].value, -math.log10(0.5))
                finally:
                    workbook.close()
            finally:
                session.halt()
                session.join(5)
                led.halt()
                led.join(5)


if __name__ == "__main__":
    unittest.main()
