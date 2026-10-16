"""Durable per-cycle files first; combined CSV/Excel on a separate worker."""
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
import csv
import json
import os
from pathlib import Path
import re
from datetime import datetime, timezone
import threading
import time
from models import Capture
from spectrometer import Settings, Spectrum


DEFAULT_RUN_CACHE = Path(__file__).resolve().parent / "run-cache"


def validate_experiment_name(name):
    if not name or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_. -]{0,79}", name):
        raise ValueError("Experiment name: 1–80 letters/numbers/spaces/._-, beginning with a letter or number.")


def write_absorbance_excel(path, wavelengths, records):
    """One sheet only; records are immutable (cycle, absorbance) snapshots."""
    from openpyxl import Workbook
    if len(records)+1 > 16384:
        raise ValueError("Excel supports at most 16,383 absorbance spectra plus wavelength.")
    path = Path(path)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Absorbance")
    sheet.append(["wavelength_nm"] + ["time_0_no_irradiation_absorbance_AU" if cycle == 0 else f"cycle_{cycle:06d}_absorbance_AU" for cycle, _ in records])
    for i, wavelength in enumerate(wavelengths):
        sheet.append([wavelength] + [values[i] for _, values in records])
    temporary = path.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary)
        with temporary.open("rb+") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def copy_file_atomically(source, target):
    target = Path(target)
    temporary = target.with_suffix(target.suffix+".tmp")
    try:
        data = Path(source).read_bytes()
        # Box can briefly lock either file during syncing. Retry only access
        # errors, on the export worker, with a bounded 1.5-second total wait.
        for attempt in range(5):
            try:
                with temporary.open("wb") as stream:
                    stream.write(data)
                    stream.flush()
                    os.fsync(stream.fileno())
                os.replace(temporary, target)
                break
            except PermissionError as exc:
                if attempt == 4:
                    raise PermissionError(
                        f"Cannot update {target}. Close this workbook in Excel and "
                        "check Box syncing and folder permissions, then use "
                        "Re-export a saved/interrupted run. "
                        f"Local workbook remains at {source}. Original error: {exc}"
                    ) from exc
                time.sleep(0.1 * 2**attempt)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except PermissionError:
            # A sync lock on cleanup must not hide the actual export error.
            pass


def recover_absorbance_excel(directory):
    """Accept a selected run folder or its private recovery folder; no hardware."""
    directory = Path(directory)
    cache = directory if (directory/"manifest.json").is_file() else DEFAULT_RUN_CACHE/directory.name
    if not (cache/"manifest.json").is_file():
        raise ValueError("No recovery records found. Choose this run's folder under python-app/run-cache.")
    manifest = json.loads((cache/"manifest.json").read_text(encoding="utf-8"))
    target_directory = Path(manifest.get("storage", {}).get("selected_output", directory))
    target_directory.mkdir(parents=True, exist_ok=True)
    dark = json.loads((cache/"dark.json").read_text(encoding="utf-8"))
    records = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(cache.glob("cycle_*.json"))]
    snapshots = [(r["cycle"], tuple(r["processed"]["absorbance"])) for r in records]
    local = cache/"combined_absorbance.xlsx"
    write_absorbance_excel(local, dark["spectrum"]["wavelengths"], snapshots)
    final = target_directory/local.name
    if local.resolve() != final.resolve():
        copy_file_atomically(local, final)
    return str(final)


def is_box_output(folder):
    """Recognize this installation's Box mount; do not guess other providers."""
    return Path(folder).is_relative_to(Path.home() / "Box")


def atomic_json(path, data):
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(data, stream, indent=2, allow_nan=False)
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary, path)


class RunStore:
    def __init__(self, folder, name, settings, dark, reference, metadata, on_error=None, combined_excel_only=False):
        validate_experiment_name(name)
        if combined_excel_only and settings.cycles > 16382:
            raise ValueError("Excel supports at most 16,382 irradiation cycles plus the time-zero measurement per run.")
        self.combined_excel_only = combined_excel_only
        self.export_records = []
        self.export_wavelengths = tuple(dark.spectrum.wavelengths)
        self.copy_label = "Combined Excel export" if combined_excel_only else "Box copying"
        folder = Path(folder).resolve()
        if not folder.is_dir():
            raise ValueError("Select an existing output folder.")
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
        self.directory = folder / f"{name}_{stamp}"
        self.directory.mkdir()
        self.cache_enabled = combined_excel_only or is_box_output(folder)
        self.working_directory = self.directory
        if self.cache_enabled:
            DEFAULT_RUN_CACHE.mkdir(parents=True, exist_ok=True)
            self.working_directory = DEFAULT_RUN_CACHE / self.directory.name
            self.working_directory.mkdir()
        self.executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="Durable run saving")
        self.mirror_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix=self.copy_label) if self.cache_enabled else None
        self.mirror_futures = []
        self.mirror_errors = []
        self.on_error = on_error
        self.closed = False
        self.manifest = {"format_version": 2, "name": name, "settings": asdict(settings), "metadata": metadata,
                         "status": "preparing", "completed_cycles": 0, "timestamp_provenance": "Host UTC/monotonic; optical timing not measured",
                         "averaging": "OceanDirect SDK scans-to-average uses selected N; no host averaging", "boxcar": "OceanDirect SDK boxcar half-width b; 2b+1; no host smoothing",
                         "intensity_provenance": "exact SDK formatted result; individual unaveraged/unsmoothed exposures unavailable",
                         "sdk_corrections": "standard OceanDirect policy: electric dark and nonlinearity enabled for preview/background/sample",
                         "timing_verified": False, "freshness_verified": False}
        atomic_json(self.working_directory / "manifest.json", self.manifest)
        atomic_json(self.working_directory / "dark.json", asdict(dark))
        atomic_json(self.working_directory / "reference.json", asdict(reference))
        self._write_capture_csv("dark", dark)
        self._write_capture_csv("reference", reference)
        if self.cache_enabled:
            self.manifest["storage"] = {"policy": "private recovery cache with asynchronous combined Excel only" if combined_excel_only else "local durable cache with asynchronous Box copies", "local_cache": str(self.working_directory), "selected_output": str(self.directory), "copies_are_not_acquisition_deadlines": True, "copy_status": "pending", "combined_excel_only": combined_excel_only}
            atomic_json(self.working_directory / "manifest.json", self.manifest)
            self._queue_mirror(["manifest.json", "dark.json", "dark.csv", "reference.json", "reference.csv"])

    def check_health(self):
        if self.mirror_errors:
            raise OSError(f"{self.copy_label} failed: {self.mirror_errors[0]}. Local run data remains in {self.working_directory}.")

    def _queue_mirror(self, names):
        if not self.cache_enabled:
            return
        self.check_health()
        self.mirror_futures = [future for future in self.mirror_futures if not future.done()]
        if len(self.mirror_futures) >= 16:
            raise OSError(f"{self.copy_label} fell behind by 16 batches. Local run data remains in {self.working_directory}; no further pulses are allowed.")
        # Immutable snapshots prevent later manifest/event updates changing a
        # queued batch. This local read is separate from slow destination IO.
        snapshots = [] if self.combined_excel_only else [(name, (self.working_directory / name).read_bytes()) for name in names if (self.working_directory / name).exists()]
        records = tuple(self.export_records) if self.combined_excel_only else ()
        def copy():
            if self.combined_excel_only:
                local = self.working_directory/"combined_absorbance.xlsx"
                write_absorbance_excel(local, self.export_wavelengths, records)
                copy_file_atomically(local, self.directory/local.name)
                return
            for name, data in snapshots:
                target = self.directory / name
                temporary = target.with_suffix(target.suffix + ".tmp")
                with temporary.open("wb") as stream:
                    stream.write(data)
                    stream.flush()
                    os.fsync(stream.fileno())
                os.replace(temporary, target)
        future = self.mirror_executor.submit(copy)
        self.mirror_futures.append(future)
        def completed(result):
            error = result.exception()
            if error:
                self.mirror_errors.append(str(error) or type(error).__name__)
                if self.on_error and len(self.mirror_errors) == 1:
                    self.on_error(f"{self.copy_label} failed: {self.mirror_errors[0]}. Close any open output workbook and check the destination. Local run data remains in {self.working_directory}.")
        future.add_done_callback(completed)

    def _drain_mirror(self):
        for future in self.mirror_futures:
            try:
                future.result()
            except Exception as exc:
                if not self.mirror_errors:
                    self.mirror_errors.append(str(exc) or type(exc).__name__)
        self.mirror_futures = []
        self.check_health()

    def _write_capture_csv(self, name, capture, processed=None):
        path = self.working_directory / f"{name}.csv"
        temporary = path.with_suffix(".csv.tmp")
        with temporary.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            sdk_result = capture.provenance.get("raw_source_scans_available") is False
            writer.writerow(["wavelength_nm", "sdk_intensity_counts" if sdk_result else "raw_mean_counts"] + ([] if sdk_result else [f"raw_scan_{i+1}_counts" for i in range(len(capture.scans))]) + (["absorbance_AU", "flag"] if processed else []))
            for i, (w, mean) in enumerate(zip(capture.spectrum.wavelengths, capture.spectrum.intensities)):
                writer.writerow([w, mean] + ([] if sdk_result else [s[i] for s in capture.scans]) + ([processed.absorbance[i], processed.flags[i]] if processed else []))
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)

    def event(self, record):
        if not self.closed:
            self.executor.submit(self._event, record)

    def _event(self, record):
        with (self.working_directory / "events.jsonl").open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(record, allow_nan=False) + "\n")
            stream.flush()

    def save_cycle(self, cycle, capture, processed, timing):
        def write():
            self.check_health()
            started = time.monotonic()
            name = f"cycle_{cycle:06d}"
            # JSON is the recovery source even if a subsequent CSV write fails.
            atomic_json(self.working_directory / f"{name}.json", {"cycle": cycle, "capture": asdict(capture), "processed": asdict(processed), "timing": timing})
            # Excel-only recovery and export both use JSON. Avoid an unused
            # CSV and its extra disk flush on the next-pulse deadline path.
            if not self.combined_excel_only:
                self._write_capture_csv(name, capture, processed)
            self.manifest["completed_cycles"] = cycle
            self.manifest["status"] = "running"
            atomic_json(self.working_directory / "manifest.json", self.manifest)
            self._event({"event": "cycle_saved_durably", "cycle": cycle, "host_save_elapsed_ms": (time.monotonic()-started)*1000, "directory": str(self.working_directory), "cached_before_box_copy": self.cache_enabled, "provenance": "host file-save timing; optical timing not measured"})
            if self.combined_excel_only:
                self.export_records.append((cycle, tuple(processed.absorbance)))
            self._queue_mirror([name + ".json", name + ".csv", "manifest.json"])
        return self.executor.submit(write)

    def finish(self, status, error=None):
        def write():
            self.manifest.update(status=status, error=error, finished_utc=datetime.now(timezone.utc).isoformat())
            atomic_json(self.working_directory / "manifest.json", self.manifest)
        return self.executor.submit(write)

    def export(self, excel=True):
        def export_and_copy():
            # Export after LED shutdown. Recovery always uses the local cache,
            # even when a destination copy fails or has not caught up yet.
            if not self.combined_excel_only:
                combined_export(self.working_directory, excel)
            if self.cache_enabled:
                try:
                    self._drain_mirror()
                    self._queue_mirror([p.name for p in self.working_directory.iterdir() if p.is_file() and p.suffix != ".tmp"])
                    self._drain_mirror()
                    self.manifest["storage"]["copy_status"] = "complete"
                    atomic_json(self.working_directory / "manifest.json", self.manifest)
                    if not self.combined_excel_only:
                        self._queue_mirror(["manifest.json"])
                        self._drain_mirror()
                except Exception as exc:
                    self.manifest["storage"].update(copy_status="error", copy_error=str(exc))
                    atomic_json(self.working_directory / "manifest.json", self.manifest)
                    raise
            return str(self.directory)
        return self.executor.submit(export_and_copy)

    def close(self):
        self.closed = True
        self.executor.shutdown(wait=True)
        if self.mirror_executor:
            self.mirror_executor.shutdown(wait=True)


def combined_export(directory, excel=True):
    """Can recover a completed/interrupted run without opening any hardware."""
    directory = Path(directory)
    records = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(directory.glob("cycle_*.json"))]
    dark = json.loads((directory / "dark.json").read_text(encoding="utf-8"))
    ref = json.loads((directory / "reference.json").read_text(encoding="utf-8"))
    wavelengths = dark["spectrum"]["wavelengths"]
    combined = directory / "combined_absorbance.csv"
    temporary = combined.with_suffix(".csv.tmp")
    headings = ["wavelength_nm"] + ["time_0_no_irradiation_absorbance_AU" if r["cycle"] == 0 else f"cycle_{r['cycle']:06d}_absorbance_AU" for r in records]
    with temporary.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headings)
        for i, w in enumerate(wavelengths):
            writer.writerow([w] + [r["processed"]["absorbance"][i] for r in records])
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary, combined)
    if excel:
        from openpyxl import Workbook
        if len(records) + 1 > 16384:
            raise ValueError("Excel's column limit exceeded; combined CSV is preserved.")
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Absorbance")
        sheet.append(headings)
        for i, w in enumerate(wavelengths):
            sheet.append([w] + [r["processed"]["absorbance"][i] for r in records])
        saved_manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        raw = workbook.create_sheet("SDK intensities" if saved_manifest.get("format_version",1)>=2 else "Raw means")
        raw.append(["wavelength_nm", "dark_counts", "reference_counts"] + [f"cycle_{r['cycle']:06d}_counts" for r in records])
        for i, w in enumerate(wavelengths):
            raw.append([w, dark["spectrum"]["intensities"][i], ref["spectrum"]["intensities"][i]] + [r["capture"]["spectrum"]["intensities"][i] for r in records])
        flags = workbook.create_sheet("Flags")
        flags.append(headings)
        for i, w in enumerate(wavelengths):
            flags.append([w] + [r["processed"]["flags"][i] for r in records])
        metadata = workbook.create_sheet("Settings and timestamps")
        metadata.append(["key", "json_value"])
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        for key, value in manifest.items():
            metadata.append([key, json.dumps(value)])
        # Avoid Excel's 32767-character cell limit for acquisition provenance.
        for r in records:
            metadata.append([f"cycle_{r['cycle']}_timing", json.dumps(r["timing"])])
            provenance = r["capture"]["provenance"]
            for index, call in enumerate(provenance.get("scan_calls", [])):
                metadata.append([f"cycle_{r['cycle']}_scan_{index+1}_host_call", json.dumps(call)])
        final = directory / "combined.xlsx"
        temporary = directory / "combined.tmp.xlsx"
        workbook.save(temporary)
        os.replace(temporary, final)
    return str(combined)
